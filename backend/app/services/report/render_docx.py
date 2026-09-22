"""Render a report doc to a .docx file (python-docx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def _fmt(v, digits=4):
    if v is None:
        return "—"
    if isinstance(v, str):
        return v
    if isinstance(v, int):
        return str(v)
    return f"{v:.{digits}g}"


def _table(doc, columns, rows):
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(columns):
        t.rows[0].cells[i].text = str(c)
    for row in rows:
        cells = t.add_row().cells
        for i, cell in enumerate(row):
            cells[i].text = "—" if cell is None else str(cell)


def render_docx(doc: dict, out_path: Path) -> Path:
    d = Document()
    d.add_heading(doc["title"], level=0)
    d.add_paragraph(f"Generated {doc['generated_at']}").italic = True

    d.add_heading("1. Project", level=1)
    d.add_paragraph().add_run("Aim. ").bold = True
    d.paragraphs[-1].add_run(doc["project"]["aim"])
    if doc["project"]["objectives"]:
        d.add_paragraph("Objectives").bold = True
        for i, o in enumerate(doc["project"]["objectives"], 1):
            d.add_paragraph(f"{i}. {o}", style="List Number")
    for label, key in [("Research questions", "research_questions"), ("Hypotheses", "hypotheses")]:
        if doc["project"][key]:
            d.add_paragraph(label).bold = True
            for x in doc["project"][key]:
                d.add_paragraph(x, style="List Bullet")

    d.add_heading("2. Dataset", level=1)
    ds = doc["dataset"]
    d.add_paragraph(
        f"{ds['filename']}"
        + (f" (sheet: {ds['sheet_name']})" if ds["sheet_name"] else "")
        + f" — {ds['n_rows']} rows × {ds['n_columns']} columns."
    )
    _table(d, ["Column", "Type", "Missing %", "Unique"],
           [[c["name"], c["dtype"], c["missing_pct"], c["unique_count"]] for c in ds["columns"]])

    d.add_heading("3. Data Quality", level=1)
    s = doc["data_quality"]["score"]
    d.add_paragraph(
        f"Overall score: {s.get('overall')}/100 (completeness {s.get('completeness')}, "
        f"validity {s.get('validity')}, consistency {s.get('consistency')}, "
        f"duplicates {s.get('duplicates')}, outliers {s.get('outliers')})."
    )
    if doc["data_quality"]["duplicate_row_count"]:
        d.add_paragraph(f"{doc['data_quality']['duplicate_row_count']} duplicate row(s) detected (not removed).")

    d.add_heading("4. Statistical Methods", level=1)
    for m in doc["methods"]:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(f"{m['label']} — ").bold = True
        p.add_run(m["description"] + (
            " (reports a Bayes factor)" if m["bayesian_supported"] else " (frequentist only in this toolset)"
        ))

    d.add_heading("5. Results", level=1)
    for idx, a in enumerate(doc["analyses"], 1):
        heading = f"{idx}. {a['method_label']}"
        if a["objective"]:
            heading += f" — {a['objective']}"
        d.add_heading(heading, level=2)
        if a["status"] != "complete":
            d.add_paragraph(f"This analysis did not complete: {a['error']}")
            continue
        d.add_paragraph(
            "Variables: " + "; ".join(f"{k} = {v}" for k, v in a["variables"].items())
            + f" · n = {a['n_used']}" + (f" ({a['n_excluded']} excluded)" if a["n_excluded"] else "")
        ).italic = True

        f = a["frequentist"]
        pf = d.add_paragraph()
        pf.add_run("Frequentist. ").bold = True
        bits = [f"{f['statistic']['name']} = {_fmt(f['statistic']['value'])}"]
        if f.get("df") is not None:
            bits.append(f"df = {f['df']}")
        if f.get("p_value") is not None:
            bits.append(f"p = {_fmt(f['p_value'])}")
        if f.get("effect_size") and f["effect_size"].get("value") is not None:
            es = f["effect_size"]
            bits.append(f"{es['name']} = {_fmt(es['value'], 3)}"
                        + (f" ({es['magnitude']})" if es.get("magnitude") else ""))
        pf.add_run("; ".join(bits) + ". " + (f.get("summary") or ""))

        b = a["bayesian"]
        pb = d.add_paragraph()
        pb.add_run("Bayesian. ").bold = True
        if b.get("available") and b.get("bayes_factor_10") is not None:
            pb.add_run(
                f"BF10 = {_fmt(b['bayes_factor_10'], 3)} — {b.get('interpretation')}. "
                "A Bayes factor is not a p-value and does not denote 'significance'."
            )
        else:
            pb.add_run(b.get("note") or "No Bayes factor available for this test.")

        for t in a["tables"]:
            d.add_paragraph(t["title"]).italic = True
            _table(d, t["columns"], t["rows"])

        failed = [c for c in a["assumptions"] if c.get("passed") is False]
        if failed:
            d.add_paragraph("Assumption checks not met:").italic = True
            for c in failed:
                d.add_paragraph(
                    f"{c['label']} — {c['detail']}"
                    + (f" {c['recommendation']}" if c.get("recommendation") else ""),
                    style="List Bullet",
                )

        chart = a.get("chart")
        if chart and Path(chart["stored_path"]).exists():
            d.add_picture(chart["stored_path"], width=Inches(5.8))

        if a["interpretation"]:
            pi = d.add_paragraph()
            pi.add_run("Interpretation. ").bold = True
            pi.add_run(a["interpretation"])

    d.add_heading("6. Limitations", level=1)
    for lim in doc["limitations"] or ["No specific limitations flagged automatically."]:
        d.add_paragraph(lim, style="List Bullet")

    d.add_heading("7. Recommendations", level=1)
    for rec in doc["recommendations"]:
        d.add_paragraph(rec, style="List Bullet")

    for section in d.sections:
        section.left_margin = section.right_margin = Inches(1)
    d.styles["Normal"].font.size = Pt(10.5)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(out_path)
    return out_path
